"""Validated tabular import and replicate-aware 4PL analysis for N3XORA."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
from collections import defaultdict
from statistics import mean
from typing import Any

import numpy as np
from scipy.optimize import least_squares
from scipy.special import expit

ANALYTICS_MODEL = "four_parameter_logistic"
ANALYTICS_VERSION = "1.0.0"
MAX_IMPORT_ROWS = 5000

ALIASES = {
    "sample_role": {"sample_role", "role", "sample_type", "type_echantillon"},
    "sample_label": {"sample_label", "sample", "sample_name", "compound", "echantillon"},
    "replicate_id": {"replicate_id", "replicate", "rep", "replicate_name", "replicat"},
    "replicate_kind": {"replicate_kind", "replicate_type", "replicate_class", "type_replicat"},
    "biological_sample_id": {"biological_sample_id", "bio_sample_id", "biological_sample", "parent_biological_replicate", "parent_replicate_id"},
    "concentration_value": {"concentration_value", "concentration", "dose", "dose_value", "conc"},
    "concentration_unit": {"concentration_unit", "concentration_uom", "dose_unit", "conc_unit"},
    "response_name": {"response_name", "readout", "assay_readout", "parameter"},
    "response_value": {"response_value", "response", "value", "measured_value"},
    "response_unit": {"response_unit", "readout_unit", "unit"},
    "qualifier": {"qualifier", "censoring", "measurement_qualifier"},
    "timepoint_value": {"timepoint_value", "timepoint", "time"},
    "timepoint_unit": {"timepoint_unit", "time_unit"},
    "well_id": {"well_id", "well", "well_position"},
    "record_key": {"record_key", "row_id", "measurement_id"},
}
REQUIRED = {
    "sample_role", "sample_label", "replicate_id", "concentration_value",
    "concentration_unit", "response_name", "response_value", "response_unit",
}
SAMPLE_ROLES = {
    "test_sample", "positive_control", "negative_control", "vehicle_control",
    "blank", "reference_standard", "untreated_control",
}


class TabularImportError(ValueError):
    """The uploaded table cannot be interpreted without guessing."""


def _header_key(value: Any) -> str:
    value = str(value or "").replace("\ufeff", "").strip().lower()
    value = re.sub(r"[\s./-]+", "_", value)
    return re.sub(r"_+", "_", value).strip("_")


def _read_table(filename: str, payload: bytes) -> tuple[list[Any], list[list[Any]]]:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise TabularImportError("Le lecteur XLSX n'est pas installé.") from exc
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            header = list(next(iterator, ()) or ())
            rows = [list(row) for row in iterator]
            workbook.close()
            return header, rows
        except Exception as exc:
            raise TabularImportError(f"Classeur XLSX illisible : {exc}") from exc
    if extension not in {"csv", "tsv", "txt"}:
        raise TabularImportError("Format non pris en charge. Utiliser CSV, TSV ou XLSX.")
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise TabularImportError("Le fichier doit être encodé en UTF-8 (BOM accepté).") from exc
    if not text.strip():
        raise TabularImportError("Le fichier est vide.")
    if extension == "tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\\t").delimiter
        except csv.Error:
            delimiter = ","
    parsed = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return (list(parsed[0]), parsed[1:]) if parsed else ([], [])


def _number(value: Any, *, required: bool, field: str) -> float | None:
    if value is None or str(value).strip() == "":
        if required:
            raise ValueError(f"{field} est obligatoire.")
        return None
    if isinstance(value, bool):
        raise ValueError(f"{field} doit être numérique.")
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        raw = str(value).strip().replace(" ", "")
        if "," in raw and "." not in raw:
            raw = raw.replace(",", ".")
        try:
            number = float(raw)
        except ValueError as exc:
            raise ValueError(f"{field} doit être numérique.") from exc
    if not math.isfinite(number):
        raise ValueError(f"{field} doit être une valeur finie.")
    return number


def parse_measurement_table(filename: str, payload: bytes) -> dict[str, Any]:
    header, rows = _read_table(filename, payload)
    source_headers = [_header_key(x) for x in header]
    if not any(source_headers):
        raise TabularImportError("La première ligne doit contenir les noms de colonnes.")
    if len(source_headers) != len(set(source_headers)):
        raise TabularImportError("La première ligne contient des noms de colonnes dupliqués.")
    column_map: dict[str, int] = {}
    for canonical, aliases in ALIASES.items():
        matches = [i for i, key in enumerate(source_headers) if key in aliases]
        if len(matches) > 1:
            raise TabularImportError(f"Plusieurs colonnes correspondent au champ {canonical}.")
        if matches:
            column_map[canonical] = matches[0]
    missing = sorted(REQUIRED - set(column_map))
    if missing:
        raise TabularImportError("Colonnes obligatoires manquantes : " + ", ".join(missing))
    if len(rows) > MAX_IMPORT_ROWS:
        raise TabularImportError(f"Le fichier dépasse la limite de {MAX_IMPORT_ROWS} lignes.")

    parsed, errors = [], []
    for line_number, row in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        if len(row) > len(header):
            errors.append({"row": line_number, "message": "La ligne contient plus de cellules que l'en-tête."})
            continue

        def cell(name: str) -> Any:
            index = column_map.get(name)
            return row[index] if index is not None and index < len(row) else None

        try:
            item = {
                "sample_role": str(cell("sample_role") or "").strip().lower().replace(" ", "_").replace("-", "_"),
                "sample_label": str(cell("sample_label") or "").strip(),
                "replicate_id": str(cell("replicate_id") or "").strip(),
                "replicate_kind": str(cell("replicate_kind") or "biological").strip().lower(),
                "biological_sample_id": str(cell("biological_sample_id") or "").strip() or None,
                "concentration_value": _number(cell("concentration_value"), required=False, field="concentration_value"),
                "concentration_unit": str(cell("concentration_unit") or "").strip() or None,
                "response_name": str(cell("response_name") or "").strip(),
                "response_value": _number(cell("response_value"), required=True, field="response_value"),
                "response_unit": str(cell("response_unit") or "").strip(),
                "qualifier": str(cell("qualifier") or "exact").strip().lower(),
                "timepoint_value": _number(cell("timepoint_value"), required=False, field="timepoint_value"),
                "timepoint_unit": str(cell("timepoint_unit") or "").strip() or None,
                "well_id": str(cell("well_id") or "").strip() or None,
                "record_key": str(cell("record_key") or "").strip() or None,
                "metadata": {"import_source": filename, "source_row": line_number},
            }
            for field in ("sample_label", "replicate_id", "response_name", "response_unit"):
                if not item[field]:
                    raise ValueError(f"{field} est obligatoire.")
            if item["concentration_value"] is not None and not item["concentration_unit"]:
                raise ValueError("Une concentration doit déclarer son unité.")
            if item["timepoint_value"] is not None and not item["timepoint_unit"]:
                raise ValueError("Un temps de mesure doit déclarer son unité.")
            if item["sample_role"] not in SAMPLE_ROLES:
                raise ValueError("sample_role non reconnu.")
            if item["replicate_kind"] not in {"biological", "technical", "independent_run"}:
                raise ValueError("replicate_kind doit être biological, technical ou independent_run.")
            if item["replicate_kind"] == "technical" and not item["biological_sample_id"]:
                raise ValueError("Chaque réplicat technique doit référencer biological_sample_id.")
            if item["replicate_kind"] != "technical" and item["biological_sample_id"] is not None:
                raise ValueError("biological_sample_id ne s’applique qu’aux réplicats techniques.")
            if item["qualifier"] not in {"exact", "below", "below_or_equal", "above", "above_or_equal"}:
                raise ValueError("qualifier non reconnu.")
            parsed.append(item)
        except (TypeError, ValueError) as exc:
            errors.append({"row": line_number, "message": str(exc)})

    return {
        "columns": sorted(column_map),
        "row_count": len(parsed),
        "measurements": parsed,
        "errors": errors[:100],
        "error_count": len(errors),
        "preview": parsed[:15],
        "source_sha256": hashlib.sha256(payload).hexdigest(),
        "source_filename": filename,
    }


def _fit_parameters(log_concentration: np.ndarray, response: np.ndarray, direction: str) -> dict[str, float]:
    if log_concentration.size < 5 or np.unique(log_concentration).size < 5:
        raise ValueError("Au moins cinq concentrations distinctes sont nécessaires pour l'ajustement 4PL.")
    if not np.all(np.isfinite(log_concentration)) or not np.all(np.isfinite(response)):
        raise ValueError("Les observations contiennent des valeurs non finies.")
    y_min, y_max = float(np.min(response)), float(np.max(response))
    span = max(y_max - y_min, float(np.std(response)), 1.0)
    lower = np.array([y_min - 10 * span, math.log(max(span * 1e-5, 1e-8)),
                      float(np.min(log_concentration)) - 3, math.log(0.05)])
    upper = np.array([y_max + 10 * span, math.log(max(span * 100, 1e-6)),
                      float(np.max(log_concentration)) + 3, math.log(10.0)])
    initial = np.array([y_min, math.log(span), float(np.median(log_concentration)), 0.0])
    initial = np.minimum(np.maximum(initial, lower + 1e-8), upper - 1e-8)
    sign = 1.0 if direction == "increasing" else -1.0

    def predict(parameters: np.ndarray) -> np.ndarray:
        bottom, log_amplitude, log_ec50, log_hill = parameters
        return bottom + math.exp(float(log_amplitude)) * expit(
            sign * math.exp(float(log_hill)) * (log_concentration - log_ec50)
        )

    fitted = least_squares(
        lambda parameters: predict(parameters) - response, x0=initial,
        bounds=(lower, upper), method="trf", loss="linear", max_nfev=5000,
    )
    if not fitted.success or not np.all(np.isfinite(fitted.x)):
        raise ValueError("L'ajustement 4PL n'a pas convergé : " + str(fitted.message))
    bottom, log_amplitude, log_ec50, log_hill = (float(x) for x in fitted.x)
    return {
        "bottom": bottom, "top": bottom + math.exp(log_amplitude),
        "log10_ec50": log_ec50, "ec50": 10.0 ** log_ec50,
        "hill_slope": math.exp(log_hill), "cost": float(fitted.cost),
    }


def fit_dose_response(
    measurements: list[Any], *, response_name: str, direction: str,
    endpoint_name: str = "IC50", timepoint_value: float | None = None,
    timepoint_unit: str | None = None, bootstrap_iterations: int = 300,
    random_seed: int = 20261031,
) -> dict[str, Any]:
    """Fit 4PL to biological-replicate means; bootstrap complete biological curves."""
    if direction not in {"increasing", "decreasing"}:
        raise ValueError("La direction doit être increasing ou decreasing.")
    if endpoint_name not in {"IC50", "EC50", "GI50", "CC50"}:
        raise ValueError("Endpoint dose-réponse non pris en charge par le modèle 4PL v1.")
    if not 100 <= int(bootstrap_iterations) <= 1000:
        raise ValueError("bootstrap_iterations doit être compris entre 100 et 1000.")
    if not response_name.strip():
        raise ValueError("Le nom du readout est obligatoire.")

    selected = [row for row in measurements
                if row.sample_role == "test_sample" and row.response_name == response_name]
    if timepoint_value is not None:
        selected = [row for row in selected if row.timepoint_value is not None
                    and float(row.timepoint_value) == float(timepoint_value)
                    and row.timepoint_unit == timepoint_unit]
    elif any(row.timepoint_value is not None for row in selected):
        raise ValueError("Des temps de mesure sont présents. Sélectionner explicitement un timepoint.")
    if not selected:
        raise ValueError("Aucune mesure ne correspond au readout et au timepoint demandés.")
    if any(row.qualifier != "exact" for row in selected):
        raise ValueError("Des mesures censurées sont présentes. Le moteur 4PL v1 exige des observations exactes.")
    if any(row.concentration_value is None or float(row.concentration_value) <= 0 for row in selected):
        raise ValueError("Toutes les mesures test_sample doivent avoir une concentration strictement positive.")
    concentration_units = {row.concentration_unit for row in selected}
    response_units = {row.response_unit for row in selected}
    if len(concentration_units) != 1 or None in concentration_units:
        raise ValueError("Plusieurs unités de concentration détectées. Aucune conversion automatique.")
    if len(response_units) != 1 or None in response_units:
        raise ValueError("Plusieurs unités de réponse détectées.")
    concentration_unit, response_unit = next(iter(concentration_units)), next(iter(response_units))

    biological = [row for row in selected if row.replicate_kind == "biological"]
    if not biological:
        raise ValueError("Aucun réplicat biologique disponible; les mesures techniques ne sont pas des unités biologiques indépendantes.")
    profiles: dict[str, dict[float, list[float]]] = defaultdict(lambda: defaultdict(list))
    for row in biological:
        concentration, response = float(row.concentration_value), float(row.response_value)
        if not math.isfinite(concentration) or not math.isfinite(response):
            raise ValueError("Les données contiennent une valeur non finie.")
        profiles[str(row.replicate_id)][concentration].append(response)

    concentrations = sorted({c for profile in profiles.values() for c in profile})
    if len(concentrations) < 5:
        raise ValueError("Au moins cinq concentrations distinctes sont requises.")
    under_replicated = [c for c in concentrations if sum(c in p for p in profiles.values()) < 2]
    if under_replicated:
        raise ValueError("Chaque concentration doit avoir au moins deux réplicats biologiques.")
    replicate_matrix = {rep: {c: float(mean(v)) for c, v in profile.items()}
                        for rep, profile in profiles.items()}
    complete = [rep for rep, profile in replicate_matrix.items()
                if all(c in profile for c in concentrations)]
    x, y = [], []
    for rep in sorted(replicate_matrix):
        for concentration in concentrations:
            if concentration in replicate_matrix[rep]:
                x.append(concentration)
                y.append(replicate_matrix[rep][concentration])
    x_array, y_array = np.asarray(x, dtype=float), np.asarray(y, dtype=float)
    log_x = np.log10(x_array)
    fitted = _fit_parameters(log_x, y_array, direction)
    sign = 1.0 if direction == "increasing" else -1.0
    predicted = fitted["bottom"] + (fitted["top"] - fitted["bottom"]) * expit(
        sign * fitted["hill_slope"] * (log_x - fitted["log10_ec50"])
    )
    residuals = y_array - predicted
    ss_res = float(np.sum(residuals ** 2))
    ss_tot = float(np.sum((y_array - float(np.mean(y_array))) ** 2))
    r_squared = 1.0 - ss_res / ss_tot if ss_tot > 0 else None
    rmse = float(math.sqrt(ss_res / len(y_array)))
    estimate = float(fitted["ec50"])
    min_c, max_c = min(concentrations), max(concentrations)
    in_range = min_c <= estimate <= max_c

    bootstrap_values = []
    if len(complete) >= 3 and in_range:
        rng = np.random.default_rng(random_seed)
        for _ in range(int(bootstrap_iterations)):
            sample = rng.choice(complete, size=len(complete), replace=True)
            bx, by = [], []
            for replicate in sample:
                for concentration in concentrations:
                    bx.append(concentration)
                    by.append(replicate_matrix[str(replicate)][concentration])
            try:
                fit = _fit_parameters(np.log10(np.asarray(bx)), np.asarray(by), direction)
                if math.isfinite(float(fit["ec50"])):
                    bootstrap_values.append(float(fit["ec50"]))
            except (ValueError, OverflowError, FloatingPointError):
                continue
    minimum_successes = max(100, int(0.70 * int(bootstrap_iterations)))
    if len(bootstrap_values) >= minimum_successes:
        ci_low, ci_high = (float(v) for v in np.percentile(bootstrap_values, [2.5, 97.5]))
        ci = {"status": "available", "level": 0.95,
              "method": "percentile_cluster_bootstrap_biological_replicates",
              "lower": ci_low, "upper": ci_high,
              "successful_resamples": len(bootstrap_values),
              "requested_resamples": int(bootstrap_iterations)}
    else:
        if not in_range:
            reason = "L'endpoint estimé est hors de la plage expérimentale : extrapolation."
        elif len(complete) < 3:
            reason = "Trois courbes biologiques complètes indépendantes au minimum sont nécessaires."
        else:
            reason = "Trop peu de réajustements bootstrap ont convergé."
        ci = {"status": "not_estimable", "level": 0.95,
              "method": "percentile_cluster_bootstrap_biological_replicates",
              "reason": reason, "successful_resamples": len(bootstrap_values),
              "requested_resamples": int(bootstrap_iterations)}

    concentration_summary = []
    for concentration in concentrations:
        values = [profile[concentration] for profile in replicate_matrix.values()
                  if concentration in profile]
        concentration_summary.append({
            "concentration": float(concentration), "concentration_unit": concentration_unit,
            "biological_replicates": len(values), "mean_response": float(np.mean(values)),
            "sd_between_biological_replicates": float(np.std(values, ddof=1)) if len(values) > 1 else None,
            "response_unit": response_unit,
        })
    grid = np.geomspace(min_c, max_c, 100)
    curve_y = fitted["bottom"] + (fitted["top"] - fitted["bottom"]) * expit(
        sign * fitted["hill_slope"] * (np.log10(grid) - fitted["log10_ec50"])
    )
    warnings = []
    if not in_range:
        warnings.append("L'endpoint est hors plage des concentrations testées; l'estimation est extrapolée.")
    if len(complete) < 3:
        warnings.append("Nombre de courbes biologiques complètes insuffisant pour un IC à 95 % par bootstrap.")
    if r_squared is not None and r_squared < 0.8:
        warnings.append("R² faible : examiner les résidus et l'adéquation du modèle.")
    if fitted["hill_slope"] <= 0.1 or fitted["hill_slope"] >= 9.9:
        warnings.append("La pente de Hill est proche d'une borne de recherche.")
    input_payload = [{
        "id": getattr(row, "id", None), "replicate_id": str(row.replicate_id),
        "concentration": str(row.concentration_value), "concentration_unit": row.concentration_unit,
        "response_name": row.response_name, "response": str(row.response_value),
        "response_unit": row.response_unit, "replicate_kind": row.replicate_kind,
        "qualifier": row.qualifier, "timepoint": str(row.timepoint_value) if row.timepoint_value is not None else None,
        "timepoint_unit": row.timepoint_unit,
    } for row in sorted(biological, key=lambda item: (str(item.replicate_id),
            float(item.concentration_value), getattr(item, "id", 0) or 0))]
    input_hash = hashlib.sha256(json.dumps(input_payload, sort_keys=True,
        separators=(",", ":")).encode("utf-8")).hexdigest()

    return {
        "algorithm": ANALYTICS_MODEL, "algorithm_version": ANALYTICS_VERSION,
        "endpoint_name": endpoint_name, "estimate": estimate,
        "concentration_unit": concentration_unit, "response_name": response_name,
        "response_unit": response_unit, "response_direction": direction,
        "hill_slope": float(fitted["hill_slope"]),
        "lower_asymptote": float(fitted["bottom"]), "upper_asymptote": float(fitted["top"]),
        "confidence_interval": ci,
        "fit_statistics": {"r_squared": float(r_squared) if r_squared is not None else None,
                           "rmse": rmse, "sum_squared_residuals": ss_res,
                           "observations_used": int(len(y_array))},
        "data_summary": {"distinct_concentrations": len(concentrations),
                         "biological_replicates_total": len(replicate_matrix),
                         "complete_biological_replicates": len(complete),
                         "non_biological_rows_excluded": len(selected) - len(biological),
                         "concentration_min": float(min_c), "concentration_max": float(max_c),
                         "endpoint_within_tested_range": bool(in_range)},
        "concentration_summary": concentration_summary,
        "fitted_curve": [{"concentration": float(cx), "predicted_response": float(cy)}
                         for cx, cy in zip(grid, curve_y)],
        "warnings": warnings, "input_sha256": input_hash,
        "interpretation_limit": (
            "Estimation exploratoire conditionnelle au modèle 4PL et au protocole déclaré. "
            "Elle ne démontre ni efficacité thérapeutique, ni causalité, ni validité clinique. "
            "Les réplicats techniques ne sont pas des unités biologiques indépendantes."
        ),
    }
